"""
Инструменты для работы с DOCX — описания ошибок, акты, пояснительные записки КМД.
"""

import os
import sys


def read_docx(file_path: str) -> str:
    """Извлечь весь текст из DOCX."""
    from docx import Document
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def read_docx_with_tables(file_path: str) -> dict:
    """Извлечь текст и таблицы из DOCX."""
    from docx import Document
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append({"table_index": i, "rows": rows})
    return {"paragraphs": paragraphs, "tables": tables}


def create_error_report(errors: list, output_path: str, title: str = "Отчёт об ошибках КМД"):
    """Создать DOCX отчёт об ошибках.

    errors: список словарей с ключами: номер, описание, тип, рекомендация
    """
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Заголовок
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Таблица ошибок
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["№", "Описание ошибки", "Тип", "Рекомендация"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for err in errors:
        row = table.add_row()
        row.cells[0].text = str(err.get("номер", ""))
        row.cells[1].text = str(err.get("описание", ""))
        row.cells[2].text = str(err.get("тип", ""))
        row.cells[3].text = str(err.get("рекомендация", ""))

    doc.save(output_path)
    print(f"Отчёт сохранён: {output_path}")


def compare_docx(file1: str, file2: str) -> dict:
    """Сравнить два DOCX файла."""
    import difflib
    text1 = read_docx(file1).splitlines()
    text2 = read_docx(file2).splitlines()
    diff = list(difflib.unified_diff(text1, text2, fromfile=file1, tofile=file2, lineterm=""))
    return {
        "added": [l for l in diff if l.startswith("+") and not l.startswith("+++")],
        "removed": [l for l in diff if l.startswith("-") and not l.startswith("---")],
        "diff": diff,
    }


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Использование: python docx_tools.py <команда> <файл>")
        print("Команды: read, tables, compare")
        sys.exit(1)

    cmd = sys.argv[1]
    if cmd == "read":
        print(read_docx(sys.argv[2]))
    elif cmd == "tables":
        import json
        result = read_docx_with_tables(sys.argv[2])
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif cmd == "compare":
        result = compare_docx(sys.argv[2], sys.argv[3])
        print(f"Добавлено: {len(result['added'])}, Удалено: {len(result['removed'])}")
        for line in result["diff"][:30]:
            print(line)
